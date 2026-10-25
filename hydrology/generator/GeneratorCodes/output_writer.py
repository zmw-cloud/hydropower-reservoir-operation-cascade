"""Write platform-compatible inflow files and auditable QC reports."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .io_utils import MonthlyHistory
from .joint_monthly import GeneratedScenario, JointMonthlyModel
from .statistics import MONTH_NAMES, QCBundle


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_value(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(_json_value(payload), ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )


def _style_excel(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 24
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in sheet.columns:
            values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
            width = min(max(max(map(len, values), default=0) + 2, 10), 42)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (float, np.floating)):
                    cell.number_format = "0.000000"
                if cell.is_date:
                    cell.number_format = "yyyy-mm-dd"
    workbook.save(path)


def station_matrix_frame(
    values: np.ndarray,
    station_index: int,
    start_year: int,
) -> pd.DataFrame:
    years = np.arange(start_year, start_year + values.shape[0], dtype=int)
    frame = pd.DataFrame(values[:, :, station_index], columns=list(MONTH_NAMES))
    frame.insert(0, "Year", years)
    return frame


def joint_long_frame(
    values: np.ndarray,
    station_names: tuple[str, ...],
    source_columns: tuple[str, ...],
    start_year: int,
) -> pd.DataFrame:
    n_years = values.shape[0]
    dates = pd.date_range(
        f"{start_year:04d}-01-01",
        periods=n_years * 12,
        freq="MS",
    )
    frame = pd.DataFrame(
        {
            "Date": dates,
            "SimulationYear": np.repeat(np.arange(1, n_years + 1), 12),
            "Year": dates.year,
            "Month": dates.month,
        }
    )
    for station_index, source_column in enumerate(source_columns):
        frame[source_column] = values[:, :, station_index].reshape(-1)
    return frame


def _write_single_sheet_excel(
    frame: pd.DataFrame,
    path: Path,
    sheet_name: str,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=sheet_name, index=False)
    _style_excel(path)


def _write_qc_excel(bundle: QCBundle, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        bundle.summary.to_excel(writer, sheet_name="Summary", index=False)
        bundle.spatial_correlations.to_excel(
            writer, sheet_name="SpatialCorrelation", index=False
        )
        bundle.monthly_statistics.to_excel(
            writer, sheet_name="MonthlyStatistics", index=False
        )
        bundle.cross_lag_correlations.to_excel(
            writer, sheet_name="CrossLag", index=False
        )
        bundle.autocorrelations.to_excel(
            writer, sheet_name="Autocorrelation", index=False
        )
        bundle.adjacent_month_correlations.to_excel(
            writer, sheet_name="AdjacentMonths", index=False
        )
        bundle.historical_log_correlation.to_excel(
            writer, sheet_name="HistoricalLogCorr", index=True
        )
        bundle.model_log_correlation.to_excel(
            writer, sheet_name="ModelLogCorr", index=True
        )
        bundle.synthetic_log_correlation.to_excel(
            writer, sheet_name="SyntheticLogCorr", index=True
        )
    _style_excel(path)


def write_scenario_outputs(
    output_root: Path,
    scenario_name: str,
    seed: int,
    start_year: int,
    history: MonthlyHistory,
    model: JointMonthlyModel,
    generated: GeneratedScenario,
    qc: QCBundle,
    mean_change_pct: np.ndarray,
    sd_change_pct: np.ndarray,
) -> list[Path]:
    """Write all requested files for one scenario and return their paths."""

    written: list[Path] = []
    station_directory_names = {
        "xiaowan": "xiaowan",
        "nuozhadu_interval": "nuozhadu_interval",
    }
    for station_index, station_name in enumerate(model.station_names):
        directory = station_directory_names.get(station_name, station_name)
        path = output_root / directory / f"SynMonthlyInflow_{scenario_name}.xlsx"
        _write_single_sheet_excel(
            station_matrix_frame(generated.values, station_index, start_year),
            path,
            sheet_name="MonthlyInflow",
        )
        written.append(path)

    joint_frame = joint_long_frame(
        generated.values,
        model.station_names,
        history.source_columns,
        start_year,
    )
    joint_csv = output_root / "joint" / f"inflow_joint_monthly_{scenario_name}.csv"
    joint_xlsx = output_root / "joint" / f"inflow_joint_monthly_{scenario_name}.xlsx"
    joint_csv.parent.mkdir(parents=True, exist_ok=True)
    joint_frame.to_csv(joint_csv, index=False, encoding="utf-8-sig", float_format="%.10f")
    _write_single_sheet_excel(joint_frame, joint_xlsx, sheet_name="JointMonthlyInflow")
    written.extend([joint_csv, joint_xlsx])

    reports_dir = output_root / "reports"
    report_json = reports_dir / f"quality_report_{scenario_name}.json"
    report_xlsx = reports_dir / f"quality_report_{scenario_name}.xlsx"
    write_json(report_json, qc.report)
    _write_qc_excel(qc, report_xlsx)
    written.extend([report_json, report_xlsx])

    sampled_years = model.historical_years[
        generated.sampled_historical_year_indices
    ]
    sampled_frame = pd.DataFrame(
        sampled_years,
        columns=[MONTH_NAMES[int(month)] for month in model.component_month_order],
    )
    sampled_frame.insert(0, "RotatedBlock", np.arange(1, len(sampled_frame) + 1))
    sampled_path = (
        reports_dir / f"sampled_historical_years_{scenario_name}.csv"
    )
    sampled_frame.to_csv(sampled_path, index=False, encoding="utf-8-sig")
    written.append(sampled_path)

    metadata = {
        "scenario": scenario_name,
        "random_seed": seed,
        "synthetic_years": int(generated.values.shape[0]),
        "calendar_start_year": start_year,
        "units": "m3/s monthly mean flow",
        "input_file": str(history.source_file),
        "input_sha256": sha256_file(history.source_file),
        "input_sheet": history.sheet_name,
        "historical_period": (
            f"{history.frame['Date'].min():%Y-%m} to "
            f"{history.frame['Date'].max():%Y-%m}"
        ),
        "station_columns": dict(zip(model.station_names, history.source_columns)),
        "requested_mean_change_pct": mean_change_pct,
        "requested_sd_change_pct": sd_change_pct,
        "method": (
            "Paired monthly historical resampling followed by regularized joint "
            "24-dimensional log-correlation reconstruction in automatically rotated "
            "overlapping annual blocks and positive exact-moment calibration."
        ),
        "cascade_note": (
            "The generator outputs Xiaowan natural inflow and Nuozhadu interval "
            "natural inflow only. During later dispatch, Nuozhadu actual inflow is "
            "Xiaowan outflow plus Nuozhadu interval inflow."
        ),
        "qc_passed": qc.report["qc_passed"],
    }
    metadata_path = reports_dir / f"metadata_{scenario_name}.json"
    write_json(metadata_path, metadata)
    written.append(metadata_path)
    return written
